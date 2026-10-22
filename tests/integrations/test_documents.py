import os
from datetime import UTC, datetime, timedelta
from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfWriter

from jarvis.integrations.documents import (
    CONTENT_SCAN_LIMIT,
    CSV_MAX_ROWS,
    DocumentError,
    DocumentSection,
    DocumentText,
    DocumentTooLargeError,
    UnsupportedDocumentError,
    default_document_roots,
    extract_text,
    find_documents,
    recent_documents,
    summarise_for_llm,
)

LONG_LINE = "The quick brown fox jumps over the lazy dog and keeps running through the field. "


def write_docx(path: Path, *, headings: bool = True) -> None:
    document = Document()
    document.add_heading("Hiring Plan", level=1)
    document.add_paragraph("We need two engineers this quarter.")
    if headings:
        document.add_heading("Budget", level=2)
    document.add_paragraph("Total budget is 300k.")
    document.add_paragraph("")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Role"
    table.rows[0].cells[1].text = "Count"
    document.save(str(path))


def write_pdf(path: Path, pages: int = 2) -> None:
    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=200, height=200)
    with path.open("wb") as handle:
        writer.write(handle)


def set_mtime(path: Path, when: datetime) -> None:
    stamp = when.timestamp()
    os.utime(path, (stamp, stamp))


@pytest.fixture
def library(tmp_path: Path) -> Path:
    """A fake Documents folder with files of varying age and type."""
    root = tmp_path / "Documents"
    (root / "projects" / "node_modules").mkdir(parents=True)
    (root / ".hidden").mkdir()
    now = datetime.now(UTC)
    files = {
        "hiring plan.txt": ("Interview loop for backend engineers.", now - timedelta(days=1)),
        "notes.md": ("# Standup\nTalked about the hiring plan.", now - timedelta(hours=2)),
        "projects/roadmap.csv": ("quarter,goal\nQ1,ship\n", now - timedelta(days=40)),
        "projects/node_modules/skip.md": ("# hiring plan in deps", now),
        ".hidden/secret plan.txt": ("hiring", now),
        "archive.zip": ("not a document", now),
        "~$hiring plan.docx": ("lock file", now),
    }
    for name, (content, modified) in files.items():
        path = root / name
        path.write_text(content, encoding="utf-8")
        set_mtime(path, modified)
    return root


# --- search ------------------------------------------------------------------


def test_default_document_roots_are_the_usual_folders(tmp_path: Path) -> None:
    assert default_document_roots(tmp_path) == [
        tmp_path / "Documents",
        tmp_path / "Desktop",
        tmp_path / "Downloads",
    ]
    assert default_document_roots()[0].parent == Path.home()


async def test_find_by_name_matches_all_words_case_insensitively(library: Path) -> None:
    matches = await find_documents([library], "PLAN hiring")

    assert [match.path.name for match in matches] == ["hiring plan.txt"]
    assert matches[0].size_bytes > 0
    assert matches[0].modified.tzinfo is UTC


async def test_find_skips_hidden_tooling_and_lock_files(library: Path) -> None:
    matches = await find_documents([library], "")

    names = {match.path.name for match in matches}
    assert names == {"hiring plan.txt", "notes.md", "roadmap.csv"}


async def test_find_ranks_newest_first_and_honours_limit(library: Path) -> None:
    all_matches = await find_documents([library], "")
    assert [match.path.name for match in all_matches] == [
        "notes.md",
        "hiring plan.txt",
        "roadmap.csv",
    ]

    assert len(await find_documents([library], "", limit=2)) == 2


async def test_find_filters_by_extension_and_age(library: Path) -> None:
    csv_only = await find_documents([library], "", extensions=[".CSV"])
    assert [match.path.name for match in csv_only] == ["roadmap.csv"]

    recent = await find_documents([library], "", modified_within=timedelta(days=7))
    assert {match.path.name for match in recent} == {"notes.md", "hiring plan.txt"}


async def test_find_searches_content_when_asked(library: Path) -> None:
    by_name = await find_documents([library], "hiring")
    assert [match.path.name for match in by_name] == ["hiring plan.txt"]

    by_content = await find_documents([library], "hiring", search_content=True)
    assert [match.path.name for match in by_content] == ["notes.md", "hiring plan.txt"]


async def test_find_content_scan_is_bounded_and_skips_unreadable(tmp_path: Path) -> None:
    for index in range(CONTENT_SCAN_LIMIT + 5):
        path = tmp_path / f"file{index:03}.txt"
        path.write_text("needle" if index == 0 else "hay", encoding="utf-8")
        set_mtime(path, datetime.now(UTC) - timedelta(minutes=index))
    (tmp_path / "broken.pdf").write_bytes(b"not a pdf")

    matches = await find_documents([tmp_path], "needle", search_content=True)

    assert [match.path.name for match in matches] == ["file000.txt"]
    too_old = tmp_path / "file999.txt"
    too_old.write_text("needle", encoding="utf-8")
    set_mtime(too_old, datetime.now(UTC) - timedelta(days=1))
    assert len(await find_documents([tmp_path], "needle", search_content=True)) == 1


async def test_missing_roots_are_ignored(library: Path, tmp_path: Path) -> None:
    matches = await find_documents([tmp_path / "nope", library], "notes")

    assert [match.path.name for match in matches] == ["notes.md"]


async def test_recent_documents_returns_newest_first(library: Path) -> None:
    recent = await recent_documents([library], limit=2)

    assert [match.path.name for match in recent] == ["notes.md", "hiring plan.txt"]


# --- extraction --------------------------------------------------------------


async def test_extract_plain_text_normalises_line_endings(tmp_path: Path) -> None:
    path = tmp_path / "a.txt"
    path.write_bytes("\ufeffline one\r\n\r\n\r\n\r\nline two  \r\n".encode())

    doc = await extract_text(path)

    assert doc == DocumentText(
        path=path, sections=(DocumentSection(None, "line one\n\nline two"),), truncated=False
    )
    assert doc.text == "line one\n\nline two"


async def test_extract_markdown_splits_on_headings(tmp_path: Path) -> None:
    path = tmp_path / "notes.md"
    path.write_text(
        "\n\nintro line\n# Title #\n\nFirst section.\n## Sub\nSecond.\n\n\nMore.\n",
        encoding="utf-8",
    )

    doc = await extract_text(path)

    assert doc.sections == (
        DocumentSection(None, "intro line"),
        DocumentSection("Title", "First section."),
        DocumentSection("Sub", "Second.\n\nMore."),
    )
    assert doc.text.startswith("intro line\n\nTitle\nFirst section.")


async def test_extract_csv_keeps_first_rows_only(tmp_path: Path) -> None:
    path = tmp_path / "data.csv"
    path.write_text("\n".join(f"{i}, name {i}" for i in range(CSV_MAX_ROWS + 10)), encoding="utf-8")

    doc = await extract_text(path)

    assert doc.truncated
    assert doc.sections[0].title == f"First {CSV_MAX_ROWS} rows"
    assert doc.sections[0].text.splitlines()[1] == "1, name 1"

    (tmp_path / "small.csv").write_text("a,b\n1,2\n", encoding="utf-8")
    small = await extract_text(tmp_path / "small.csv")
    assert not small.truncated
    assert small.sections[0] == DocumentSection("2 rows", "a, b\n1, 2")


async def test_extract_docx_groups_by_headings_and_appends_tables(tmp_path: Path) -> None:
    path = tmp_path / "plan.docx"
    write_docx(path)

    doc = await extract_text(path)

    assert doc.sections == (
        DocumentSection("Hiring Plan", "We need two engineers this quarter."),
        DocumentSection("Budget", "Total budget is 300k."),
        DocumentSection("Tables", "Role | Count"),
    )


async def test_extract_docx_without_headings_is_one_section(tmp_path: Path) -> None:
    path = tmp_path / "flat.docx"
    write_docx(path, headings=False)

    doc = await extract_text(path)

    assert doc.sections[0].title == "Hiring Plan"
    assert doc.sections[0].text == "We need two engineers this quarter.\nTotal budget is 300k."


async def test_extract_docx_rejects_non_docx_bytes(tmp_path: Path) -> None:
    path = tmp_path / "fake.docx"
    path.write_bytes(b"PK garbage")

    with pytest.raises(DocumentError, match="not a valid Word document"):
        await extract_text(path)


async def test_extract_pdf_yields_one_section_per_page(tmp_path: Path) -> None:
    path = tmp_path / "blank.pdf"
    write_pdf(path, pages=3)

    doc = await extract_text(path)

    assert [section.title for section in doc.sections] == ["Page 1", "Page 2", "Page 3"]
    assert not doc.truncated


async def test_extract_pdf_rejects_malformed_and_encrypted_files(tmp_path: Path) -> None:
    broken = tmp_path / "broken.pdf"
    broken.write_bytes(b"%PDF-1.4 but nothing else")
    with pytest.raises(DocumentError, match="malformed PDF"):
        await extract_text(broken)

    encrypted = tmp_path / "locked.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=10, height=10)
    writer.encrypt("secret")
    with encrypted.open("wb") as handle:
        writer.write(handle)
    with pytest.raises(DocumentError, match="encrypted"):
        await extract_text(encrypted)


async def test_extract_rejects_unsupported_missing_and_oversized_files(tmp_path: Path) -> None:
    with pytest.raises(UnsupportedDocumentError, match=r"unsupported file type \.xlsx"):
        await extract_text(tmp_path / "sheet.xlsx")
    with pytest.raises(UnsupportedDocumentError, match=r"\(none\)"):
        await extract_text(tmp_path / "noext")
    with pytest.raises(DocumentError, match="file not found"):
        await extract_text(tmp_path / "missing.txt")

    big = tmp_path / "big.txt"
    big.write_text("x" * 100, encoding="utf-8")
    with pytest.raises(DocumentTooLargeError, match="exceeds") as info:
        await extract_text(big, max_bytes=10)
    assert info.value.path == big


async def test_extract_caps_characters_across_sections(tmp_path: Path) -> None:
    path = tmp_path / "long.md"
    path.write_text(f"# A\n{LONG_LINE * 3}\n# B\n{LONG_LINE * 3}\n# C\nnever", encoding="utf-8")

    doc = await extract_text(path, max_chars=300)

    assert doc.truncated
    assert [section.title for section in doc.sections] == ["A", "B"]
    assert doc.sections[1].text.endswith("…")
    assert sum(len(section.text) for section in doc.sections) <= 300


# --- summarise_for_llm -------------------------------------------------------


def make_doc(*sections: tuple[str | None, str], truncated: bool = False) -> DocumentText:
    return DocumentText(
        path=Path("C:/docs/report.docx"),
        sections=tuple(DocumentSection(title, text) for title, text in sections),
        truncated=truncated,
    )


def test_summarise_keeps_everything_when_it_fits() -> None:
    doc = make_doc(("Intro", "short"), (None, "no heading"))

    assert summarise_for_llm(doc, 1000) == "# report.docx\n\n## Intro\n\nshort\n\nno heading"


def test_summarise_trims_long_sections_but_keeps_short_ones_and_headings() -> None:
    doc = make_doc(("Summary", "Key point."), ("Chapter", LONG_LINE * 20), ("End", "Done."))

    rendered = summarise_for_llm(doc, 400)

    assert len(rendered) <= 400
    assert "## Summary\n\nKey point." in rendered
    assert "## Chapter" in rendered
    assert "## End\n\nDone." in rendered
    assert "…" in rendered
    assert "\n\n…" not in rendered


def test_summarise_marks_truncated_documents_and_cuts_on_word_boundaries() -> None:
    doc = make_doc((None, "alpha beta gamma delta"), truncated=True)

    rendered = summarise_for_llm(doc, 40)

    assert rendered.startswith("# report.docx (truncated)")
    assert rendered.endswith("alpha beta…")


def test_summarise_never_exceeds_budget_even_when_headings_alone_do() -> None:
    doc = make_doc(*[(f"Heading number {i}", "body") for i in range(20)])

    assert len(summarise_for_llm(doc, 60)) <= 60
