from collections.abc import Callable
from pathlib import Path

import pytest
from docx import Document

from jarvis.agent.confirmation import DECLINED_SPOKEN, ConfirmationPolicy
from jarvis.agent.registry import ToolRegistry
from jarvis.agent.tool import ToolContext
from jarvis.config import Settings, default_documents_output_dir
from jarvis.integrations.services import IntegrationServices
from jarvis.integrations.stubs import StubEmailProvider
from jarvis.tools.draft_document import DraftDocumentInput, DraftDocumentTool, file_stem
from tests.agent.fakes import FakeConfirmation
from tests.tools.integration_helpers import integration_context

CONTENT = """# Design review brief

The review moves to **Thursday 14:00** because the venue is booked.

## Actions
- Book room 3B
- Tell **Tom**
1. Send agenda
2. Prepare demo

```
uv run pytest
```
"""


@pytest.fixture
def settings(make_settings: Callable[..., Settings], tmp_path: Path) -> Settings:
    return make_settings(documents_output_dir=tmp_path / "out")


@pytest.fixture
def context(settings: Settings) -> ToolContext:
    return integration_context(settings, IntegrationServices(email=(StubEmailProvider(),)))


def test_is_discovered_and_confirms_only_overwrites(context: ToolContext) -> None:
    tool = ToolRegistry.discover().get("draft_document")

    assert isinstance(tool, DraftDocumentTool)
    assert not tool.requires_confirmation
    assert not tool.needs_confirmation(DraftDocumentInput(title="New", content="x"), context)


def test_output_dir_setting(make_settings: Callable[..., Settings]) -> None:
    assert make_settings().documents_output_dir == default_documents_output_dir()
    assert make_settings(documents_output_dir="").documents_output_dir == (
        Path.home() / "Documents" / "Jarvis"
    )
    assert make_settings(documents_output_dir="x/y").documents_output_dir == Path("x/y")


@pytest.mark.parametrize(
    ("title", "filename", "expected"),
    [
        ("Q3 plan: draft!", None, "Q3-plan-draft"),
        ("???", None, "document"),
        ("ignored", "My notes v2", "My notes v2"),
    ],
)
def test_file_stem(title: str, filename: str | None, expected: str) -> None:
    assert file_stem(DraftDocumentInput(title=title, content="x", filename=filename)) == expected


async def test_writes_markdown_with_a_source_citation(
    context: ToolContext, settings: Settings
) -> None:
    result = await DraftDocumentTool().call(
        {
            "title": "Review brief",
            "content": "Short body.",
            "source_email_id": "stub-email:stub-mail-1",
        },
        context,
    )

    path = settings.documents_output_dir / "Review-brief.md"
    assert result.ok
    assert result.spoken == "I've saved Review-brief.md in your documents folder."
    assert path.read_text(encoding="utf-8") == (
        "# Review brief\n\n"
        '> Source: email from Maya Lindqvist <maya@example.com>, "Design review moved to '
        'Thursday", received 2026-10-21 08:18\n\n'
        "Short body.\n"
    )
    assert result.data is not None
    assert result.data["overwritten"] is False
    assert list(settings.documents_output_dir.iterdir()) == [path]  # no temp file left


async def test_content_with_its_own_title_is_not_given_a_second(
    context: ToolContext, settings: Settings
) -> None:
    await DraftDocumentTool().call({"title": "Brief", "content": CONTENT}, context)

    text = (settings.documents_output_dir / "Brief.md").read_text(encoding="utf-8")
    assert text.startswith("# Design review brief\n")


async def test_writes_a_word_document_with_styles(context: ToolContext, settings: Settings) -> None:
    result = await DraftDocumentTool().call(
        {"title": "Brief", "content": CONTENT, "format": "docx", "filename": "brief"}, context
    )

    assert result.ok
    document = Document(str(settings.documents_output_dir / "brief.docx"))
    styled = [(p.style.name if p.style else "", p.text) for p in document.paragraphs]
    assert styled == [
        ("Heading 1", "Design review brief"),
        ("Normal", "The review moves to Thursday 14:00 because the venue is booked."),
        ("Heading 2", "Actions"),
        ("List Bullet", "Book room 3B"),
        ("List Bullet", "Tell Tom"),
        ("List Number", "Send agenda"),
        ("List Number", "Prepare demo"),
        ("No Spacing", "uv run pytest"),
    ]
    bold = [run.text for run in document.paragraphs[1].runs if run.bold]
    assert bold == ["Thursday 14:00"]
    assert document.core_properties.title == "Brief"


async def test_word_document_gets_title_and_source(
    context: ToolContext, settings: Settings
) -> None:
    await DraftDocumentTool().call(
        {
            "title": "Notes",
            "content": "Line one\nline two\n\nNext paragraph",
            "format": "docx",
            "source_email_id": "stub-email:stub-mail-2",
        },
        context,
    )

    document = Document(str(settings.documents_output_dir / "Notes.docx"))
    texts = [p.text for p in document.paragraphs]
    assert texts[0] == "Notes"
    assert texts[1].startswith("Source: email from GitHub")
    assert texts[2:] == ["Line one line two", "Next paragraph"]
    assert document.core_properties.comments.startswith("Source: email from GitHub")


async def test_overwrite_asks_and_respects_no(settings: Settings) -> None:
    target = settings.documents_output_dir / "Brief.md"
    target.parent.mkdir(parents=True)
    target.write_text("original", encoding="utf-8")
    asker = FakeConfirmation(answers=[False, True])
    context = integration_context(settings, IntegrationServices(), confirmation=asker)
    arguments = {"title": "Brief", "content": "new"}

    declined = await ConfirmationPolicy().execute_with_policy(
        DraftDocumentTool(), arguments, context
    )
    assert declined.spoken == DECLINED_SPOKEN
    assert target.read_text(encoding="utf-8") == "original"

    accepted = await ConfirmationPolicy().execute_with_policy(
        DraftDocumentTool(), arguments, context
    )

    assert (
        asker.prompts
        == ["Brief.md already exists in your documents folder. Do you want me to overwrite it?"] * 2
    )
    assert accepted.ok
    assert accepted.data is not None
    assert accepted.data["overwritten"] is True
    assert target.read_text(encoding="utf-8") == "# Brief\n\nnew\n"


async def test_source_email_needs_the_email_integration(settings: Settings) -> None:
    context = integration_context(settings, None)

    result = await DraftDocumentTool().call(
        {"title": "x", "content": "y", "source_email_id": "gmail:1"}, context
    )

    assert not result.ok
    assert "isn't configured: email" in result.spoken


@pytest.mark.parametrize("filename", ["../escape", "sub/dir", "C:evil", ".hidden", "a\\b"])
async def test_filenames_cannot_leave_the_output_folder(
    context: ToolContext, filename: str
) -> None:
    result = await DraftDocumentTool().call(
        {"title": "x", "content": "y", "filename": filename}, context
    )

    assert not result.ok
    assert "filename" in result.detail


async def test_a_failed_render_leaves_no_temp_file(
    context: ToolContext, settings: Settings, monkeypatch: pytest.MonkeyPatch
) -> None:
    def explode(*_args: object) -> str:
        raise OSError("disk full")

    monkeypatch.setattr("jarvis.tools.draft_document.render_markdown", explode)

    result = await DraftDocumentTool().call({"title": "x", "content": "y"}, context)

    assert not result.ok
    assert "disk full" in result.detail
    assert list(settings.documents_output_dir.iterdir()) == []
