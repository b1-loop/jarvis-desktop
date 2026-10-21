"""Local document search and text extraction.

Backs "find my notes on the hiring plan", "open the last thing I worked on"
and "summarise that PDF". Search walks the configured roots (by default the
user's Documents, Desktop and Downloads) and ranks by modification time,
because the file the user means is almost always a recent one. Extraction
turns PDF, Word, Markdown, plain text and CSV into sections of text with
hard size caps, and :func:`summarise_for_llm` trims that to a context budget
without losing the headings that tell the model what it is looking at.

Every filesystem and parsing call runs in ``asyncio.to_thread`` so the voice
pipeline stays responsive while a large PDF is being read.
"""

from __future__ import annotations

import asyncio
import csv
import io
import os
import re
from collections.abc import Iterator, Sequence
from dataclasses import dataclass
from datetime import UTC, datetime, timedelta
from pathlib import Path
from typing import Final

from jarvis.utils.errors import JarvisError
from jarvis.utils.logging import get_logger

log = get_logger(__name__)

DEFAULT_EXTENSIONS: Final[tuple[str, ...]] = (".pdf", ".docx", ".txt", ".md", ".csv")
DEFAULT_FIND_LIMIT: Final = 10
DEFAULT_RECENT_LIMIT: Final = 5
DEFAULT_MAX_FILE_BYTES: Final = 50 * 1024 * 1024
DEFAULT_MAX_CHARS: Final = 200_000
CSV_MAX_ROWS: Final = 50
CONTENT_SCAN_LIMIT: Final = 40
"""How many of the most recent candidates a content search opens."""

_CONTENT_SCAN_CHARS: Final = 20_000
_SKIPPED_DIRECTORIES: Final = frozenset(
    {"node_modules", "__pycache__", ".git", ".venv", "venv", ".cache", "AppData", "$RECYCLE.BIN"}
)
_DEFAULT_ROOT_NAMES: Final = ("Documents", "Desktop", "Downloads")
_MARKDOWN_HEADING: Final = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")
_ELLIPSIS: Final = "…"


class DocumentError(JarvisError):
    """A document could not be read.

    Attributes:
        path: The file in question.
    """

    def __init__(self, path: Path, reason: str) -> None:
        self.path = path
        super().__init__(f"cannot read {path}: {reason}")


class UnsupportedDocumentError(DocumentError):
    """The file extension is not one :func:`extract_text` knows."""


class DocumentTooLargeError(DocumentError):
    """The file exceeds the size cap and will not be opened."""


@dataclass(frozen=True, slots=True)
class DocumentMatch:
    """A file found by :func:`find_documents` or :func:`recent_documents`.

    Attributes:
        path: Absolute path.
        modified: Last modification time, timezone-aware UTC.
        size_bytes: File size.
    """

    path: Path
    modified: datetime
    size_bytes: int


@dataclass(frozen=True, slots=True)
class DocumentSection:
    """A page, heading-delimited chunk or other natural unit of a document.

    Attributes:
        title: Section heading or page label; ``None`` for unstructured text.
        text: The section body.
    """

    title: str | None
    text: str


@dataclass(frozen=True, slots=True)
class DocumentText:
    """Extracted text of one document.

    Attributes:
        path: Where it came from.
        sections: Natural units in order (pages for PDF, headings for docx/md).
        truncated: True if the character cap cut content.
    """

    path: Path
    sections: tuple[DocumentSection, ...]
    truncated: bool

    @property
    def text(self) -> str:
        """All sections joined, headings included as their own lines."""
        return "\n\n".join(
            f"{section.title}\n{section.text}" if section.title else section.text
            for section in self.sections
        )


def default_document_roots(home: Path | None = None) -> list[Path]:
    """The folders people keep documents in, under ``home`` (default: the user's)."""
    base = home or Path.home()
    return [base / name for name in _DEFAULT_ROOT_NAMES]


async def find_documents(
    roots: Sequence[Path],
    query: str,
    *,
    extensions: Sequence[str] = DEFAULT_EXTENSIONS,
    limit: int = DEFAULT_FIND_LIMIT,
    modified_within: timedelta | None = None,
    search_content: bool = False,
) -> list[DocumentMatch]:
    """Find documents whose name (or, optionally, content) matches ``query``.

    Every whitespace-separated word of ``query`` must appear in the file name,
    case-insensitively. With ``search_content`` the most recent
    :data:`CONTENT_SCAN_LIMIT` candidates are also opened and searched; that
    is bounded on purpose, since a user asking by content means a recent file.

    Returns:
        Matches, most recently modified first, at most ``limit``.
    """
    return await asyncio.to_thread(
        _find_documents_sync,
        roots,
        query,
        extensions=tuple(extensions),
        limit=limit,
        modified_within=modified_within,
        search_content=search_content,
    )


async def recent_documents(
    roots: Sequence[Path],
    *,
    limit: int = DEFAULT_RECENT_LIMIT,
    extensions: Sequence[str] = DEFAULT_EXTENSIONS,
) -> list[DocumentMatch]:
    """The most recently modified documents under ``roots``."""
    return await asyncio.to_thread(_recent_documents_sync, roots, tuple(extensions), limit)


def _find_documents_sync(
    roots: Sequence[Path],
    query: str,
    *,
    extensions: tuple[str, ...],
    limit: int,
    modified_within: timedelta | None,
    search_content: bool,
) -> list[DocumentMatch]:
    words = [word.lower() for word in query.split()]
    cutoff = datetime.now(UTC) - modified_within if modified_within else None
    candidates = sorted(
        (match for match in _walk(roots, extensions) if cutoff is None or match.modified >= cutoff),
        key=lambda match: match.modified,
        reverse=True,
    )
    matches = [match for match in candidates if _name_matches(match.path, words)]
    if search_content:
        matches.extend(_content_matches(candidates, words, exclude=matches))
        matches.sort(key=lambda match: match.modified, reverse=True)
    log.info("documents.find", query=query, candidates=len(candidates), matches=len(matches))
    return matches[:limit]


def _recent_documents_sync(
    roots: Sequence[Path], extensions: tuple[str, ...], limit: int
) -> list[DocumentMatch]:
    matches = sorted(_walk(roots, extensions), key=lambda match: match.modified, reverse=True)
    return matches[:limit]


def _walk(roots: Sequence[Path], extensions: tuple[str, ...]) -> Iterator[DocumentMatch]:
    """Yield every document under ``roots``, skipping hidden and tooling folders."""
    wanted = {extension.lower() for extension in extensions}
    for root in roots:
        if not root.is_dir():
            log.debug("documents.root.missing", root=str(root))
            continue
        for directory, dirnames, filenames in os.walk(root):
            dirnames[:] = [
                name
                for name in dirnames
                if not name.startswith(".") and name not in _SKIPPED_DIRECTORIES
            ]
            for filename in filenames:
                path = Path(directory) / filename
                if path.suffix.lower() in wanted and not filename.startswith("~$"):
                    yield _describe(path)


def _describe(path: Path) -> DocumentMatch:
    stat = path.stat()
    return DocumentMatch(
        path=path.resolve(),
        modified=datetime.fromtimestamp(stat.st_mtime, tz=UTC),
        size_bytes=stat.st_size,
    )


def _name_matches(path: Path, words: list[str]) -> bool:
    name = path.name.lower()
    return all(word in name for word in words)


def _content_matches(
    candidates: list[DocumentMatch], words: list[str], exclude: list[DocumentMatch]
) -> list[DocumentMatch]:
    """Open the most recent candidates and keep those whose text has every word."""
    found: list[DocumentMatch] = []
    excluded = {match.path for match in exclude}
    for match in candidates[:CONTENT_SCAN_LIMIT]:
        if match.path in excluded:
            continue
        try:
            text = _extract_text_sync(match.path, _CONTENT_SCAN_CHARS, DEFAULT_MAX_FILE_BYTES).text
        except DocumentError as error:
            log.debug("documents.content_scan.skipped", path=str(match.path), error=str(error))
            continue
        if all(word in text.lower() for word in words):
            found.append(match)
    return found


async def extract_text(
    path: Path,
    *,
    max_chars: int = DEFAULT_MAX_CHARS,
    max_bytes: int = DEFAULT_MAX_FILE_BYTES,
) -> DocumentText:
    """Extract the text of a PDF, Word, Markdown, text or CSV file.

    Args:
        path: The document.
        max_chars: Cap on extracted characters; the result is marked
            ``truncated`` when it bites.
        max_bytes: Files larger than this are refused outright.

    Raises:
        UnsupportedDocumentError: For any other extension.
        DocumentTooLargeError: If the file exceeds ``max_bytes``.
        DocumentError: If the file is missing, encrypted or malformed.
    """
    return await asyncio.to_thread(_extract_text_sync, path, max_chars, max_bytes)


def _extract_text_sync(path: Path, max_chars: int, max_bytes: int) -> DocumentText:
    suffix = path.suffix.lower()
    if suffix not in DEFAULT_EXTENSIONS:
        raise UnsupportedDocumentError(path, f"unsupported file type {suffix or '(none)'}")
    if not path.is_file():
        raise DocumentError(path, "file not found")
    size = path.stat().st_size
    if size > max_bytes:
        raise DocumentTooLargeError(path, f"{size} bytes exceeds the {max_bytes} byte cap")
    sections, truncated = _READERS[suffix](path)
    sections, capped = _cap_sections(sections, max_chars)
    log.info("documents.extracted", path=str(path), sections=len(sections), truncated=truncated)
    return DocumentText(path=path, sections=tuple(sections), truncated=truncated or capped)


def _read_pdf(path: Path) -> tuple[list[DocumentSection], bool]:
    # Lazy: pypdf is in the optional ``documents`` extra.
    from pypdf import PdfReader
    from pypdf.errors import PyPdfError

    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            raise DocumentError(path, "the PDF is encrypted")
        pages = [page.extract_text() or "" for page in reader.pages]
    except PyPdfError as error:
        raise DocumentError(path, f"malformed PDF ({error})") from error
    sections = [
        DocumentSection(title=f"Page {number}", text=_tidy(text))
        for number, text in enumerate(pages, start=1)
    ]
    return sections, False


def _read_docx(path: Path) -> tuple[list[DocumentSection], bool]:
    # Lazy: python-docx is in the optional ``documents`` extra.
    from docx import Document
    from docx.opc.exceptions import PackageNotFoundError

    try:
        document = Document(str(path))
    except PackageNotFoundError as error:
        raise DocumentError(path, "not a valid Word document") from error
    sections = _group_by_headings(
        (paragraph.style.name if paragraph.style is not None else "", paragraph.text)
        for paragraph in document.paragraphs
    )
    tables = [
        " | ".join(cell.text.strip() for cell in row.cells)
        for table in document.tables
        for row in table.rows
    ]
    if tables:
        sections.append(DocumentSection(title="Tables", text="\n".join(tables)))
    return sections, False


def _group_by_headings(paragraphs: Iterator[tuple[str, str]]) -> list[DocumentSection]:
    """Split Word paragraphs into sections at ``Heading N`` / ``Title`` styles."""
    sections: list[DocumentSection] = []
    title: str | None = None
    body: list[str] = []
    for style, text in paragraphs:
        if style.startswith(("Heading", "Title")) and text.strip():
            _append_section(sections, title, body)
            title, body = text.strip(), []
        elif text.strip():
            body.append(text.strip())
    _append_section(sections, title, body)
    return sections


def _append_section(sections: list[DocumentSection], title: str | None, body: list[str]) -> None:
    if title is not None or body:
        sections.append(DocumentSection(title=title, text="\n".join(body)))


def _read_markdown(path: Path) -> tuple[list[DocumentSection], bool]:
    sections: list[DocumentSection] = []
    title: str | None = None
    body: list[str] = []
    for line in _read_text_file(path).splitlines():
        if match := _MARKDOWN_HEADING.match(line):
            _append_section(sections, title, body)
            title, body = match.group(2), []
        else:
            body.append(line)
    _append_section(sections, title, body)
    tidied = [DocumentSection(section.title, _tidy(section.text)) for section in sections]
    return [section for section in tidied if section.title is not None or section.text], False


def _read_plain(path: Path) -> tuple[list[DocumentSection], bool]:
    return [DocumentSection(title=None, text=_tidy(_read_text_file(path)))], False


def _read_csv(path: Path) -> tuple[list[DocumentSection], bool]:
    """The first :data:`CSV_MAX_ROWS` rows as comma-joined lines."""
    reader = csv.reader(io.StringIO(_read_text_file(path)))
    rows: list[str] = []
    truncated = False
    for row in reader:
        if len(rows) >= CSV_MAX_ROWS:
            truncated = True
            break
        rows.append(", ".join(cell.strip() for cell in row))
    title = f"First {len(rows)} rows" if truncated else f"{len(rows)} rows"
    return [DocumentSection(title=title, text="\n".join(rows))], truncated


def _read_text_file(path: Path) -> str:
    """Decode as UTF-8 (BOM tolerated); undecodable bytes become U+FFFD."""
    return path.read_bytes().decode("utf-8-sig", errors="replace")


_READERS: Final = {
    ".pdf": _read_pdf,
    ".docx": _read_docx,
    ".txt": _read_plain,
    ".md": _read_markdown,
    ".csv": _read_csv,
}


def _cap_sections(
    sections: list[DocumentSection], max_chars: int
) -> tuple[list[DocumentSection], bool]:
    """Keep sections in order until the character budget runs out."""
    kept: list[DocumentSection] = []
    remaining = max_chars
    for section in sections:
        if len(section.text) <= remaining:
            kept.append(section)
            remaining -= len(section.text)
            continue
        kept.append(DocumentSection(section.title, _trim(section.text, remaining)))
        return kept, True
    return kept, False


def summarise_for_llm(doc: DocumentText, max_chars: int) -> str:
    """Render ``doc`` within ``max_chars``, trimming bodies before headings.

    Section bodies share the budget by water-filling: short sections keep all
    their text and long ones are cut to the remaining fair share, so a table
    of contents or a short conclusion is never lost to a long chapter.
    """
    header = f"# {doc.path.name}" + (" (truncated)" if doc.truncated else "")
    headings = [f"## {section.title}" if section.title else "" for section in doc.sections]
    fixed_cost = len(header) + sum(len(heading) + 4 for heading in headings)
    budgets = _water_fill([len(section.text) for section in doc.sections], max_chars - fixed_cost)
    parts = [header]
    for section, heading, budget in zip(doc.sections, headings, budgets, strict=True):
        if heading:
            parts.append(heading)
        if body := _trim(section.text, budget):
            parts.append(body)
    rendered = "\n\n".join(parts)
    # Only reachable when the headings alone exceed the budget.
    return rendered if len(rendered) <= max_chars else _trim(rendered, max_chars)


def _water_fill(lengths: list[int], budget: int) -> list[int]:
    """Allocate ``budget`` across items so small ones are fully served first."""
    allocations = [0] * len(lengths)
    remaining = max(0, budget)
    for position, index in enumerate(sorted(range(len(lengths)), key=lambda i: lengths[i])):
        share = remaining // (len(lengths) - position)
        allocations[index] = min(lengths[index], share)
        remaining -= allocations[index]
    return allocations


def _trim(text: str, budget: int) -> str:
    """Cut ``text`` to ``budget`` characters on a word boundary, marking the cut."""
    if len(text) <= budget:
        return text
    if budget <= len(_ELLIPSIS):
        return ""
    longest = budget - len(_ELLIPSIS)
    cut = text.rfind(" ", 0, longest + 1)
    return text[: cut if cut > 0 else longest].rstrip() + _ELLIPSIS


def _tidy(text: str) -> str:
    """Normalise line endings and collapse runs of blank lines."""
    lines = [line.rstrip() for line in text.replace("\r\n", "\n").split("\n")]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


# Design notes:
# - os.walk with in-place pruning instead of Path.rglob: rglob cannot skip
#   node_modules or .git subtrees, and a Downloads folder with one cloned repo
#   would dominate every search.
# - Content search is capped at the CONTENT_SCAN_LIMIT most recent files. An
#   unbounded full-text scan of Documents belongs to an index (Windows Search,
#   Everything), not to a voice assistant answering in real time.
# - Parsers are imported inside the reader functions so the module imports,
#   and plain-text extraction works, without the ``documents`` extra.
# - DocumentText keeps sections rather than one string so the LLM budgeting
#   in summarise_for_llm can keep every heading and trim bodies fairly.
